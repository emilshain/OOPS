"""Rebuild OOPS.docx experiment sections (canonical, safe to rerun).

Everything before the first "EXP n" heading (e.g. the user's own lines) is
kept untouched. From the first heading on, the document is rebuilt:

    EXP n
    <blank>
    label1.java
    code lines...
    <blank>                (separator: page break allowed here)
    label2.java            (only for experiments with several code files)
    code lines...
    <blank>                (separator: page break allowed here)
    OUTPUT:
    <blank>
    output lines...
    <blank>                (separator: page break allowed here)

Page-break rules: keepLines on every paragraph + keepNext chains glue
[EXP n + blank + code file], each further code file, and [OUTPUT: + blank +
outputs] into blocks. Every block is also chained to its trailing separator
blank, so Word only breaks at separator blanks - never inside code or inside
output (unless a single block is taller than one page, i.e. overflow).
"""
import os
from docx import Document

ROOT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(ROOT, "OOPS.docx")

EXPS = [
    ("EXP 1", ["exp1/palindrome.java"], "exp1/output.txt"),
    ("EXP 2", ["exp2/frequency.java"], "exp2/output.txt"),
    ("EXP 3", ["exp3/Matrix.java"], "exp3/output.txt"),
    ("EXP 4", ["exp4/EmployeeDemo.java"], "exp4/output.txt"),
    ("EXP 5", ["exp5/ShapeDemo.java"], "exp5/output.txt"),
    ("EXP 6", ["exp6/AccessDemo.java", "exp6/pkg1/ClassA.java", "exp6/pkg1/ClassB.java", "exp6/pkg2/ClassC.java"], "exp6/output.txt"),
    ("EXP 7", ["exp7/Excep.java"], "exp7/output.txt"),
    ("EXP 8", ["exp8/VoteDemo.java"], "exp8/output.txt"),
]


def read_lines(rel):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return f.read().replace("\r\n", "\n").rstrip("\n").split("\n")


def label(rel):
    return rel.split("/", 1)[1]


def glue(paras):
    """keepLines everywhere + keepNext on every paragraph (incl. the last,
    so the block also binds to its trailing separator blank)."""
    for p in paras:
        pf = p.paragraph_format
        pf.keep_together = True
        pf.keep_with_next = True


def separate(paras):
    """Separator blanks: page break allowed here."""
    for p in paras:
        pf = p.paragraph_format
        pf.keep_together = True
        pf.keep_with_next = False


def rebuild():
    doc = Document(DOCX)
    paras = doc.paragraphs

    first_exp = next((i for i, p in enumerate(paras) if p.text.startswith("EXP ")), None)
    if first_exp is None:
        raise SystemExit("no EXP headings found")

    # drop everything from the first EXP heading to the end
    for p in paras[first_exp:]:
        p._element.getparent().remove(p._element)

    for title, java_files, out_file in EXPS:
        # block 1: heading + blank + first code file
        b = [doc.add_paragraph(title), doc.add_paragraph(""), doc.add_paragraph(label(java_files[0]))]
        b += [doc.add_paragraph(line) for line in read_lines(java_files[0])]
        glue(b)

        for jf in java_files[1:]:
            separate([doc.add_paragraph("")])  # blank between code files
            b = [doc.add_paragraph(label(jf))]
            b += [doc.add_paragraph(line) for line in read_lines(jf)]
            glue(b)

        separate([doc.add_paragraph("")])  # blank between code and OUTPUT:

        ob = [doc.add_paragraph("OUTPUT:"), doc.add_paragraph("")]
        ob += [doc.add_paragraph(line) for line in read_lines(out_file)]
        glue(ob)

        separate([doc.add_paragraph("")])  # trailing blank before next EXP

    doc.save(DOCX)
    print("rebuilt OK:", len(doc.paragraphs), "paragraphs")


if __name__ == "__main__":
    rebuild()
