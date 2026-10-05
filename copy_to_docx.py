import os
from docx import Document

ROOT = os.path.dirname(os.path.abspath(__file__))
DOCX = os.path.join(ROOT, "OOPS.docx")

EXPS = [
    ("EXP 1", ["exp1/palindrome.java"], "exp1/output.txt"),
    ("EXP 2", ["exp2/frequency.java"], "exp2/output.txt"),
    ("EXP 3", ["exp3/Matrix.java"], "exp3/output.txt"),
    ("EXP 4", ["exp4/EmployeeDemo.java"], "exp4/output.txt"),
    ("EXP 5", ["exp5/ShapeDemo.java"], "exp5/output.txt"),
    ("EXP 6", ["exp6/AccessDemo.java", "exp6/pkg1/ClassA.java", "exp6/pkg1/ClassB.java", "exp6/pkg2/ClassC.java"], "exp6/output.txt"),
    ("EXP 7", ["exp7/Excep.java"], "exp7/output.txt"),
    ("EXP 8", ["exp8/VoteDemo.java"], "exp8/output.txt"),
]


def read_lines(rel):
    with open(os.path.join(ROOT, rel), encoding="utf-8") as f:
        return f.read().replace("\r\n", "\n").rstrip("\n").split("\n")


def label(rel):
    # path relative to the exp folder only, e.g. palindrome.java, pkg1/ClassA.java
    return rel.split("/", 1)[1]


def keep_block(paras, last_free):
    # keepLines on every paragraph + keepNext chain so Word never splits the
    # block across a page break; only the block's final paragraph may separate.
    for i, p in enumerate(paras):
        pf = p.paragraph_format
        pf.keep_together = True
        pf.keep_with_next = (i < len(paras) - 1) or not last_free


doc = Document(DOCX)

for title, java_files, out_file in EXPS:
    heading = doc.add_paragraph(title)
    code_paras = [heading]
    for jf in java_files:
        code_paras.append(doc.add_paragraph(label(jf)))
        for line in read_lines(jf):
            code_paras.append(doc.add_paragraph(line))
    out_paras = [doc.add_paragraph("OUTPUT:")]
    for line in read_lines(out_file):
        out_paras.append(doc.add_paragraph(line))
    doc.add_paragraph()

    keep_block(code_paras, last_free=False)  # code stays glued to OUTPUT:
    keep_block(out_paras, last_free=True)    # output may end a page

doc.save(DOCX)
print("Saved", DOCX)
